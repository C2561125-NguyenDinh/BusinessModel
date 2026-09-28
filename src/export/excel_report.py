"""Workbook Excel nhiều sheet có định dạng, lọc theo phạm vi. Mọi giá trị đọc từ outputs/ của pipeline."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data import Scope, apply_scope, scope_metrics, series_in_scope

INK, PETROL, TEAL, MIST, AMBER, CORAL, GREEN = "0E2A33", "12505F", "1B8A8F", "EEF5F5", "E8A33D", "D05A4E", "3C8D5A"
F_TITLE = Font(name="Calibri", size=20, bold=True, color=INK)
F_SUB = Font(name="Calibri", size=11, color="5B6B70")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FILL_HDR = PatternFill("solid", fgColor=PETROL)
FILL_BAND = PatternFill("solid", fgColor=MIST)
THIN = Border(bottom=Side(style="thin", color="D5E1E3"))
RISK_FILL = {"HIGH": "F9DEDA", "MEDIUM": "FBEBD0", "LOW": "DCEFE2"}

# (tên cột gốc, tiêu đề tiếng Việt, định dạng số)
COLS = {
    "cv": [("model", "Mô hình", None), ("WAPE_mean", "WAPE TB", "0.000%"), ("WAPE_std", "Độ lệch chuẩn WAPE", "0.000%"),
           ("MAE_mean", "MAE TB", "#,##0.000"), ("RMSE_mean", "RMSE TB", "#,##0.000"), ("sMAPE_mean", "sMAPE TB", "0.0000"),
           ("RMSLE_mean", "RMSLE TB", "0.0000"), ("train_seconds_mean", "Giây huấn luyện TB", "#,##0.0"), ("rank_WAPE", "Hạng", "0"), ("selected", "Được chọn", None)],
    "folds": [("fold", "Lần", "0"), ("model", "Mô hình", None), ("val_start", "Từ ngày", None), ("val_end", "Đến ngày", None),
              ("train_rows", "Dòng huấn luyện", "#,##0"), ("WAPE", "WAPE", "0.000%"), ("MAE", "MAE", "#,##0.000"), ("RMSE", "RMSE", "#,##0.000")],
    "holdout": [("model", "Mô hình", None), ("selected_by_cv", "Chọn bằng CV", None), ("WAPE", "WAPE", "0.000%"), ("MAE", "MAE", "#,##0.000"),
                ("RMSE", "RMSE", "#,##0.000"), ("sMAPE", "sMAPE", "0.0000"), ("RMSLE", "RMSLE", "0.0000"), ("train_rows", "Dòng huấn luyện", "#,##0")],
    "pred": [("date", "Ngày", "dd/mm/yyyy"), ("store_nbr", "Cửa hàng", "0"), ("family", "Nhóm hàng", None), ("onpromotion", "onpromotion", "#,##0"),
             ("sales", "Thực tế", "#,##0.00"), ("point", "Dự báo điểm", "#,##0.00"), ("p10", "P10", "#,##0.00"), ("p50", "P50", "#,##0.00"), ("p90", "P90", "#,##0.00")],
    "inv": [("store_nbr", "Cửa hàng", "0"), ("family", "Nhóm hàng", None), ("risk", "Nhãn rủi ro", None), ("p10", "P10 kỳ", "#,##0.0"), ("p50", "P50 kỳ", "#,##0.0"),
            ("p90", "P90 kỳ", "#,##0.0"), ("critical_fractile", "q*", "0.00"), ("assumed_inventory_position", "Vị thế giả định (0,8×P50)", "#,##0.0"),
            ("reorder_point_proxy", "Mức mục tiêu (P90)", "#,##0.0"), ("safety_stock_proxy", "Tồn kho an toàn (P90−P50)", "#,##0.0"),
            ("suggested_replenishment", "Lượng bổ sung gợi ý", "#,##0.0")],
    "queue": [("priority_rank", "Ưu tiên #", "0"), ("case_id", "Mã phiếu", None), ("store_nbr", "Cửa hàng", "0"), ("family", "Nhóm hàng", None),
              ("risk", "Nhãn", None), ("priority_score", "Điểm", "0"), ("reason_codes", "Mã lý do", None), ("p50", "P50 kỳ", "#,##0.0"),
              ("relative_uncertainty_r", "r", "0.000"), ("action", "Hành động", None), ("chosen_option", "Phương án", None),
              ("recommended_qty", "Số lượng đề xuất", "#,##0"), ("target_stock", "Mức mục tiêu", "#,##0"), ("confidence", "Tin cậy", None),
              ("promotion_stance", "Khuyến mãi", None), ("promo_extrapolation", "Ngoại suy KM", None), ("approval_status", "Trạng thái", None)],
    "err": [("store_nbr", "Cửa hàng", "0"), ("family", "Nhóm hàng", None), ("actual_sum", "Tổng thực tế", "#,##0.0"), ("pred_sum", "Tổng dự báo", "#,##0.0"),
            ("WAPE", "WAPE", "0.00%"), ("MAE", "MAE", "#,##0.00"), ("bias", "Bias (thực tế − dự báo)", "#,##0.00"),
            ("coverage_P10_P90", "Bao phủ P10–P90", "0.0%"), ("mean_interval_width", "Độ rộng TB", "#,##0.00")],
}


def _title(ws, title, sub):
    ws["A1"] = title; ws["A1"].font = F_TITLE
    ws["A2"] = sub; ws["A2"].font = F_SUB
    ws.row_dimensions[1].height = 30


def _table(ws, df, cols, r0=4, c0=1, widths=None, band=True):
    cols = [c for c in cols if c[0] in df.columns]
    for j, (_, h, _) in enumerate(cols):
        cell = ws.cell(r0, c0 + j, h)
        cell.font, cell.fill = F_HDR, FILL_HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 32
    vals = df[[c[0] for c in cols]].itertuples(index=False, name=None)
    for i, row in enumerate(vals, 1):
        for j, v in enumerate(row):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif hasattr(v, "item"):
                v = v.item()
            cell = ws.cell(r0 + i, c0 + j, v)
            fmt = cols[j][2]
            if fmt:
                cell.number_format = fmt
            if band and i % 2 == 0:
                cell.fill = FILL_BAND
            cell.border = THIN
    last_r, last_c = r0 + len(df), c0 + len(cols) - 1
    ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(last_c)}{max(last_r, r0 + 1)}"
    ws.auto_filter.ref = ref
    ws.freeze_panes = ws.cell(r0 + 1, c0)
    for j, (k, h, _) in enumerate(cols):
        w = (widths or {}).get(k) or min(max(len(h) + 2, 11, int(df[k].astype(str).str.len().quantile(.9)) + 2 if len(df) else 11), 46)
        ws.column_dimensions[get_column_letter(c0 + j)].width = w
    return {h: get_column_letter(c0 + j) for j, (_, h, _) in enumerate(cols)}, last_r


def _risk_rule(ws, col, r1, r2):
    for k, fill in RISK_FILL.items():
        ws.conditional_formatting.add(f"{col}{r1}:{col}{r2}", CellIsRule(operator="equal", formula=[f'"{k}"'], fill=PatternFill("solid", fgColor=fill)))


def build_workbook(b: dict, sc: Scope | None = None) -> bytes:
    sc = sc or Scope()
    keys = series_in_scope(b, sc)
    m = scope_metrics(b, keys)
    winner = (b.get("selection") or {}).get("winner", "XGBoost")
    stamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    wb = Workbook()

    # ---------- Bìa
    ws = wb.active; ws.title = "Bìa"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 70
    ws["B2"] = "FAVORITA · FORECAST INTELLIGENCE"; ws["B2"].font = Font(size=11, bold=True, color=TEAL)
    ws["B3"] = "Báo cáo dữ liệu dự báo & hỗ trợ quyết định"; ws["B3"].font = Font(size=24, bold=True, color=INK)
    ws["B4"] = f"Lần chạy chính thức · mô hình được chọn: {winner} · kỳ dự báo {m['period'][0]} → {m['period'][1]}"; ws["B4"].font = F_SUB
    info = [("Phạm vi", sc.label()), ("Số chuỗi trong phạm vi", m["series"]), ("Thời điểm xuất", stamp),
            ("Nguồn", "Thư mục outputs/ của pipeline (Stage 01–12). Tệp này không tạo số liệu mới."),
            ("Lưu ý", "Tồn kho là mô phỏng theo giả định (Cu = 3, Co = 1, vị thế = 0,8 × P50); kịch bản khuyến mãi là dự báo có điều kiện, không phải tác động nhân quả.")]
    for i, (k, v) in enumerate(info):
        ws.cell(6 + i, 2, k).font = Font(bold=True, color=PETROL); c = ws.cell(6 + i, 3, v); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[6 + i].height = 32 if len(str(v)) > 70 else 18
    ws.cell(12, 2, "Các sheet trong tệp").font = Font(bold=True, size=13, color=INK)
    sheets = [("KPI", "Chỉ số chính của phạm vi và toàn hệ thống, kèm biểu đồ"), ("Mô hình – CV", "Tổng hợp và từng lần kiểm định chéo cuốn chiếu"),
              ("Kiểm tra cuối", "Chỉ số 5 mô hình trên 31/07–15/08/2017"), ("Dự báo theo ngày", "Thực tế, dự báo điểm, P10/P50/P90 theo ngày × chuỗi"),
              ("Tồn kho", "Đại lượng Newsvendor và nhãn rủi ro theo chuỗi"), ("Khuyến mãi", "P50 16 ngày theo 5 kịch bản onpromotion"),
              ("Hàng đợi quyết định", "Phiếu khuyến nghị Stage 08 theo thứ tự ưu tiên"), ("Sai số theo chuỗi", "WAPE, bias, bao phủ theo chuỗi")]
    for i, (k, v) in enumerate(sheets):
        c = ws.cell(13 + i, 2, k); c.hyperlink = f"#'{k}'!A1"; c.font = Font(color=TEAL, underline="single", bold=True)
        ws.cell(13 + i, 3, v).font = F_SUB

    # ---------- KPI
    ws = wb.create_sheet("KPI"); ws.sheet_view.showGridLines = False
    _title(ws, "Chỉ số chính", sc.label())
    pm = b["prob_metrics"].iloc[0]; hw = b["holdout"].set_index("model").loc[winner]; cw = b["cv_summary"].set_index("model").loc[winner]
    rows = [("Phạm vi đang chọn", None, None),
            ("Số chuỗi", m["series"], "#,##0"), ("Tổng thực tế 16 ngày", m["actual_sum"], "#,##0"), ("Tổng dự báo điểm 16 ngày", m["point_sum"], "#,##0"),
            ("WAPE dự báo điểm", m["WAPE"], "0.000%"), ("Bao phủ P10–P90", m["coverage"], "0.00%"), ("Độ rộng P10–P90 TB (đơn vị/ngày)", m["width"], "#,##0.000"),
            ("Chuỗi HIGH / MEDIUM / LOW", f"{m['risk_counts']['HIGH']} / {m['risk_counts']['MEDIUM']} / {m['risk_counts']['LOW']}", None),
            ("Tổng lượng bổ sung gợi ý (mô phỏng)", m["replenishment_sum"], "#,##0"),
            ("Toàn hệ thống (1.782 chuỗi)", None, None),
            (f"WAPE CV {winner}", cw.WAPE_mean, "0.000%"), (f"WAPE kiểm tra cuối {winner}", hw.WAPE, "0.000%"),
            ("Pinball P10 / P50 / P90", f"{pm.Pinball_P10:.3f} / {pm.Pinball_P50:.3f} / {pm.Pinball_P90:.3f}".replace(".", ","), None),
            ("Bao phủ P10–P90 (danh nghĩa 80%)", pm.Coverage_P10_P90, "0.00%"), ("Phương pháp bất định", pm.uncertainty_method, None)]
    ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 30
    for i, (k, v, f) in enumerate(rows, 4):
        a = ws.cell(i, 1, k)
        if v is None:
            a.font = Font(bold=True, color="FFFFFF"); a.fill = FILL_HDR; ws.cell(i, 2).fill = FILL_HDR
            continue
        a.font = Font(color=INK); c = ws.cell(i, 2, v); c.font = Font(bold=True, size=12, color=PETROL)
        c.alignment = Alignment(horizontal="right", wrap_text=True)
        if f:
            c.number_format = f
        a.border = c.border = THIN
    ws.row_dimensions[18].height = 45
    # dữ liệu biểu đồ
    cvs = b["cv_summary"].sort_values("WAPE_mean")
    ws["D4"], ws["E4"] = "Mô hình", "WAPE CV (%)"
    for i, r in enumerate(cvs.itertuples(), 5):
        ws.cell(i, 4, r.model); ws.cell(i, 5, round(100 * r.WAPE_mean, 3))
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "WAPE trung bình 3 lần kiểm định (%)"
    ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=4 + len(cvs)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=4, min_row=5, max_row=4 + len(cvs)))
    ch.y_axis.majorGridlines = None; ch.legend = None; ch.height, ch.width = 7.5, 15
    ch.series[0].graphicalProperties.solidFill = TEAL; ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, "G3")
    dl = b["err_level"]
    ws["D12"], ws["E12"] = "Mức nhu cầu", "Bao phủ P10–P90 (%)"
    for i, r in enumerate(dl.itertuples(), 13):
        ws.cell(i, 4, r.demand_level); ws.cell(i, 5, round(100 * r.coverage_P10_P90, 2))
    ch = BarChart(); ch.type = "col"; ch.style = 10; ch.title = "Bao phủ P10–P90 theo mức nhu cầu (%)"
    ch.add_data(Reference(ws, min_col=5, min_row=12, max_row=12 + len(dl)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=4, min_row=13, max_row=12 + len(dl)))
    ch.legend = None; ch.height, ch.width = 7.5, 15; ch.series[0].graphicalProperties.solidFill = PETROL
    ws.add_chart(ch, "G22")
    # tổng theo ngày của phạm vi
    p = apply_scope(b["pred"], keys)
    daily = p.groupby("date")[["sales", "p10", "p50", "p90"]].sum().reset_index()
    ws["D22"], ws["E22"], ws["F22"] = "Ngày", "Thực tế", "P50"
    for i, r in enumerate(daily.itertuples(), 23):
        ws.cell(i, 4, r.date.to_pydatetime()).number_format = "dd/mm"; ws.cell(i, 5, round(r.sales, 1)); ws.cell(i, 6, round(r.p50, 1))
    lc = LineChart(); lc.title = "Tổng thực tế và P50 theo ngày – phạm vi"; lc.style = 12; lc.height, lc.width = 7.5, 15
    lc.add_data(Reference(ws, min_col=5, max_col=6, min_row=22, max_row=22 + len(daily)), titles_from_data=True)
    lc.set_categories(Reference(ws, min_col=4, min_row=23, max_row=22 + len(daily)))
    lc.series[0].graphicalProperties.line.solidFill = INK; lc.series[1].graphicalProperties.line.solidFill = TEAL
    ws.add_chart(lc, "G41")
    for col in "DEF":
        ws.column_dimensions[col].width = 14

    # ---------- Mô hình – CV
    ws = wb.create_sheet("Mô hình – CV")
    _title(ws, "So sánh 5 mô hình – kiểm định chéo cuốn chiếu", "3 lần × 16 ngày; chọn mô hình bằng WAPE trung bình trước khi mở tập kiểm tra cuối.")
    h, last = _table(ws, b["cv_summary"].sort_values("WAPE_mean"), COLS["cv"])
    ws.conditional_formatting.add(f"{h['WAPE TB']}5:{h['WAPE TB']}{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL))
    r0 = last + 3
    ws.cell(r0 - 1, 1, "Từng lần kiểm định").font = Font(bold=True, size=13, color=INK)
    _table(ws, b["cv_folds"].sort_values(["fold", "WAPE"]), COLS["folds"], r0=r0)
    ws.freeze_panes = "A5"

    # ---------- Kiểm tra cuối
    ws = wb.create_sheet("Kiểm tra cuối")
    _title(ws, "Tập kiểm tra cuối 31/07–15/08/2017", "Mô hình huấn luyện trên mọi ngày trước 31/07/2017.")
    h, last = _table(ws, b["holdout"].sort_values("WAPE"), COLS["holdout"])
    ws.conditional_formatting.add(f"{h['WAPE']}5:{h['WAPE']}{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color=AMBER))

    # ---------- Dự báo theo ngày
    ws = wb.create_sheet("Dự báo theo ngày")
    _title(ws, "Dự báo xác suất theo ngày × chuỗi", f"{len(p):,} dòng · P10/P90 = dự báo điểm + phân vị phần dư Mondrian của nhóm quy mô.".replace(",", "."))
    _table(ws, p.sort_values(["store_nbr", "family", "date"]), COLS["pred"], widths={"family": 28, "date": 12}, band=len(p) < 6000)

    # ---------- Tồn kho
    ws = wb.create_sheet("Tồn kho")
    inv = apply_scope(b["inv"], keys).sort_values("suggested_replenishment", ascending=False)
    a = b["inv_assump"].iloc[0]
    _title(ws, "Hỗ trợ quyết định tồn kho (mô phỏng)", f"Cu = {a.Cu_assumption:g}, Co = {a.Co_assumption:g}, q* = {a.q_star:g}. {a.note}")
    h, last = _table(ws, inv, COLS["inv"], widths={"family": 28})
    _risk_rule(ws, h["Nhãn rủi ro"], 5, last)
    ws.conditional_formatting.add(f"{h['Lượng bổ sung gợi ý']}5:{h['Lượng bổ sung gợi ý']}{last}", DataBarRule(start_type="min", end_type="max", color=TEAL))

    # ---------- Khuyến mãi
    ws = wb.create_sheet("Khuyến mãi")
    s = apply_scope(b["scen"], keys)
    pv = s.pivot_table(index=["store_nbr", "family"], columns="scenario", values="p50").reset_index()
    order = [c for c in ["No promotion", "Low", "Current", "Medium", "High"] if c in pv.columns]
    pv = pv[["store_nbr", "family"] + order]
    pv["delta_high"] = pv["High"] - pv["Current"]
    pv["rel_high"] = pv.apply(lambda r: (r.High / r.Current - 1) if r.Current > 0 else None, axis=1)
    _title(ws, "P50 16 ngày theo kịch bản onpromotion", "Dự báo có điều kiện của mô hình được chọn; onpromotion không phải % giảm giá; thay đổi > 100% là ngoại suy.")
    cols = [("store_nbr", "Cửa hàng", "0"), ("family", "Nhóm hàng", None)] + [(c, c, "#,##0.0") for c in order] + [
        ("delta_high", "High − Current", "#,##0.0"), ("rel_high", "High so với Current", "0.0%")]
    h, last = _table(ws, pv.sort_values("delta_high", ascending=False), cols, widths={"family": 28})
    ws.conditional_formatting.add(f"{h['High so với Current']}5:{h['High so với Current']}{last}",
                                  CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill("solid", fgColor="F9DEDA")))

    # ---------- Hàng đợi quyết định
    ws = wb.create_sheet("Hàng đợi quyết định")
    q = apply_scope(b["queue"], keys).sort_values("priority_rank")
    _title(ws, "Hàng đợi phiếu khuyến nghị (Stage 08)", "Điểm = LARGE_SERIES +3 · RISK_HIGH +2 · RISK_MEDIUM +1 · PROMO_EXTRAPOLATION +1. Mọi phiếu chờ con người phê duyệt.")
    h, last = _table(ws, q, COLS["queue"], widths={"family": 26, "reason_codes": 30, "action": 36, "case_id": 26})
    _risk_rule(ws, h["Nhãn"], 5, last)

    # ---------- Sai số theo chuỗi
    ws = wb.create_sheet("Sai số theo chuỗi")
    e = apply_scope(b["err_series"], keys).sort_values("actual_sum", ascending=False)
    _title(ws, "Sai số theo chuỗi trên tập kiểm tra cuối", "Sai số theo dự báo điểm; bias = TB(thực tế − dự báo).")
    h, last = _table(ws, e, COLS["err"], widths={"family": 28})
    ws.conditional_formatting.add(f"{h['Bao phủ P10–P90']}5:{h['Bao phủ P10–P90']}{last}",
                                  CellIsRule(operator="lessThan", formula=["0.5"], fill=PatternFill("solid", fgColor="F9DEDA")))

    for w in wb.worksheets:
        w.sheet_properties.tabColor = TEAL if w.title in ("Bìa", "KPI") else PETROL
        w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.sheet_properties.pageSetUpPr.fitToPage = True; w.page_setup.fitToHeight = 0
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()
