"""Mẫu dự báo để con người ra quyết định (Excel).

Phần thông tin dự báo lấy nguyên từ Stage 05/07/08 và được khóa; người lập kế hoạch chỉ nhập các ô vàng.
Công thức Excel tự tính lượng bổ sung theo tồn kho thực, chênh lệch so với đề xuất và các cảnh báo kiểm soát.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .data import Scope, apply_scope, series_in_scope

INK, PETROL, TEAL, MIST, AMBER = "0E2A33", "12505F", "1B8A8F", "EEF5F5", "E8A33D"
INPUT = PatternFill("solid", fgColor="FFF4CC")
LOCK = PatternFill("solid", fgColor="F3F6F7")
HDR_LOCK = PatternFill("solid", fgColor=PETROL)
HDR_IN = PatternFill("solid", fgColor="B7791F")
HDR_CALC = PatternFill("solid", fgColor=TEAL)
THIN = Border(left=Side(style="thin", color="D5E1E3"), right=Side(style="thin", color="D5E1E3"),
              top=Side(style="thin", color="D5E1E3"), bottom=Side(style="thin", color="D5E1E3"))
DECISIONS = ["Duyệt", "Điều chỉnh", "Từ chối", "Yêu cầu xem lại"]
OPTIONS = ["A", "B", "C", "Khác"]

# (khóa, tiêu đề, loại: L = khóa, I = nhập, F = công thức, định dạng, độ rộng)
LAYOUT = [
    ("priority_rank", "Ưu tiên #", "L", "0", 8), ("case_id", "Mã phiếu", "L", None, 22), ("store_nbr", "Cửa hàng", "L", "0", 9),
    ("family", "Nhóm hàng", "L", None, 24), ("risk", "Nhãn rủi ro", "L", None, 10), ("confidence", "Tin cậy", "L", None, 11),
    ("reason_codes", "Mã lý do", "L", None, 26), ("p10", "P10 kỳ", "L", "#,##0", 10), ("p50", "P50 kỳ", "L", "#,##0", 10),
    ("p90", "P90 kỳ", "L", "#,##0", 10), ("tA", "Mục tiêu A (P50)", "L", "#,##0", 11), ("tB", "Mục tiêu B (q*)", "L", "#,##0", 11),
    ("tC", "Mục tiêu C (P90)", "L", "#,##0", 11), ("chosen_option", "PA đề xuất", "L", None, 9), ("recommended_qty", "SL đề xuất (vị thế giả định)", "L", "#,##0", 13),
    ("promotion_stance", "Khuyến mãi", "L", None, 20), ("promo_extrapolation", "Ngoại suy KM", "L", None, 9),
    ("in_stock", "Tồn kho thực tế", "I", "#,##0", 12), ("in_transit", "Hàng đang về", "I", "#,##0", 11),
    ("in_decision", "Quyết định", "I", None, 15), ("in_option", "Phương án chọn", "I", None, 10),
    ("f_target", "Mục tiêu theo PA", "F", "#,##0", 12), ("f_need", "SL cần theo tồn thực", "F", "#,##0", 12),
    ("in_final", "Số lượng cuối", "I", "#,##0", 12), ("f_diff", "Chênh so với đề xuất", "F", "0.0%", 11),
    ("in_reason", "Lý do / ghi chú", "I", None, 34), ("in_reviewer", "Người duyệt", "I", None, 16), ("in_date", "Ngày duyệt", "I", "dd/mm/yyyy", 12),
    ("f_check", "Kiểm soát", "F", None, 26),
]


def _rows(b: dict, keys: pd.DataFrame, cases: list | None) -> pd.DataFrame:
    q = apply_scope(b["queue"], keys).sort_values("priority_rank").copy()
    if cases:
        opt = {c["case_id"]: {o["key"]: o["target"] for o in c["recommendation"]["options"]} for c in cases}
        for k in "ABC":
            q["t" + k] = q.case_id.map(lambda i, k=k: opt.get(i, {}).get(k))
    else:  # dự phòng: A = P50, C = P90, B nội suy như Stage 08
        q["tA"], q["tC"] = q.p50, q.p90
        q["tB"] = q.p50 + (q.p90 - q.p50) * (0.75 - 0.5) / (0.9 - 0.5)
    q["reason_codes"] = q.reason_codes.str.replace("|", " · ", regex=False)
    q["promo_extrapolation"] = q.promo_extrapolation.map({True: "Có", False: "Không"})
    return q


def build_template(b: dict, sc: Scope | None = None, cases: list | None = None, max_rows: int | None = None) -> bytes:
    sc = sc or Scope()
    keys = series_in_scope(b, sc)
    q = _rows(b, keys, cases)
    if max_rows:
        q = q.head(max_rows)
    p = apply_scope(b["pred"], keys)
    period = (p.date.min().strftime("%d/%m/%Y"), p.date.max().strftime("%d/%m/%Y")) if len(p) else ("", "")
    winner = (b.get("selection") or {}).get("winner", "XGBoost")
    wb = Workbook()

    # ------------------------------------------------ Hướng dẫn
    ws = wb.active; ws.title = "Hướng dẫn"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 30; ws.column_dimensions["C"].width = 92
    ws["B2"] = "MẪU RA QUYẾT ĐỊNH DỰ BÁO"; ws["B2"].font = Font(size=11, bold=True, color=TEAL)
    ws["B3"] = "Phiếu phê duyệt bổ sung hàng theo dự báo P10/P50/P90"; ws["B3"].font = Font(size=22, bold=True, color=INK)
    ws["B4"] = f"Kỳ dự báo 16 ngày: {period[0]} → {period[1]} · mô hình {winner} · {len(q)} chuỗi · {sc.label()}"; ws["B4"].font = Font(color="5B6B70")
    guide = [
        ("1. Đọc dự báo", "Các cột nền xám (P10, P50, P90, mục tiêu A/B/C, đề xuất) lấy nguyên từ pipeline và được khóa. P50 là trung vị; khoảng P10–P90 kỳ vọng chứa khoảng 80% khả năng (bao phủ thực tế của lần chạy: 75,40%)."),
        ("2. Nhập tồn kho thực", "Nhập 'Tồn kho thực tế' và 'Hàng đang về'. Nếu để trống, cột 'SL cần theo tồn thực' dùng số lượng đề xuất của pipeline (tính theo vị thế giả định 0,8 × P50)."),
        ("3. Chọn quyết định", "Duyệt: đồng ý đề xuất · Điều chỉnh: đổi phương án hoặc số lượng · Từ chối: không bổ sung · Yêu cầu xem lại: cần thêm thông tin."),
        ("4. Chọn phương án", "A = mức P50 (tồn dư thấp, ~50% nguy cơ thiếu hàng) · B = q* = 0,75 (cân bằng Cu/Co) · C = P90 (bảo thủ) · Khác = tự nhập số lượng."),
        ("5. Ghi số lượng & lý do", "Nhập 'Số lượng cuối', 'Lý do', 'Người duyệt', 'Ngày duyệt'. Theo POL-HITL-001, mọi quyết định phải có người duyệt và lý do."),
        ("6. Kiểm soát tự động", "Cột 'Kiểm soát' cảnh báo khi thiếu lý do/người duyệt, số lượng cuối vượt mục tiêu P90 hoặc lệch hơn 30% so với nhu cầu theo tồn thực, và khi kịch bản khuyến mãi có ngoại suy."),
        ("Giới hạn", "Chi phí Cu = 3, Co = 1 là giả định; không có giá bán, biên lợi nhuận, thời gian cung ứng thực. Kịch bản khuyến mãi là dự báo có điều kiện, không phải tác động nhân quả."),
    ]
    for i, (k, v) in enumerate(guide, 6):
        ws.cell(i, 2, k).font = Font(bold=True, color=PETROL)
        c = ws.cell(i, 3, v); c.alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[i].height = 36
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.cell(14, 2, "Chú thích màu").font = Font(bold=True, size=12)
    for i, (fill, t) in enumerate([(LOCK, "Dữ liệu dự báo – khóa"), (INPUT, "Ô người lập kế hoạch nhập"), (PatternFill("solid", fgColor="DDF1F1"), "Công thức tự tính")], 15):
        ws.cell(i, 2).fill = fill; ws.cell(i, 3, t)

    # ------------------------------------------------ Mẫu quyết định
    ws = wb.create_sheet("Mẫu quyết định")
    ws["A1"] = "PHIẾU PHÊ DUYỆT BỔ SUNG HÀNG THEO DỰ BÁO"; ws["A1"].font = Font(size=16, bold=True, color=INK)
    ws["A2"] = f"Kỳ {period[0]} → {period[1]} · {winner} · đơn vị: tổng 16 ngày · ô vàng = nhập"; ws["A2"].font = Font(color="5B6B70")
    # dải nhóm cột
    groups = [("THÔNG TIN DỰ BÁO (KHÓA)", "L", HDR_LOCK), ("NGƯỜI LẬP KẾ HOẠCH NHẬP / TỰ TÍNH", None, HDR_IN)]
    first_in = next(i for i, c in enumerate(LAYOUT, 1) if c[2] != "L")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=first_in - 1)
    ws.merge_cells(start_row=3, start_column=first_in, end_row=3, end_column=len(LAYOUT))
    for col, (t, _, f) in ((1, groups[0]), (first_in, groups[1])):
        c = ws.cell(3, col, t); c.fill = f; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(horizontal="center")
    HR = 4
    col = {}
    for j, (k, h, t, fmt, w) in enumerate(LAYOUT, 1):
        col[k] = get_column_letter(j)
        c = ws.cell(HR, j, h); c.font = Font(bold=True, color="FFFFFF")
        c.fill = {"L": HDR_LOCK, "I": HDR_IN, "F": HDR_CALC}[t]
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = THIN
        ws.column_dimensions[col[k]].width = w
    ws.row_dimensions[HR].height = 44
    n = len(q); r1, r2 = HR + 1, HR + max(n, 1)
    for i, row in enumerate(q.itertuples(index=False), r1):
        rd = row._asdict()
        for j, (k, h, t, fmt, w) in enumerate(LAYOUT, 1):
            c = ws.cell(i, j); c.border = THIN
            L = lambda key: f"{col[key]}{i}"  # noqa: E731
            if t == "L":
                v = rd.get(k); c.value = v.item() if hasattr(v, "item") else v; c.fill = LOCK
            elif t == "I":
                c.fill = INPUT; c.protection = Protection(locked=False)
            else:
                c.fill = PatternFill("solid", fgColor="DDF1F1")
                if k == "f_target":
                    c.value = (f'=IF({L("in_option")}="A",{L("tA")},IF({L("in_option")}="B",{L("tB")},IF({L("in_option")}="C",{L("tC")},'
                               f'IF({L("chosen_option")}="A",{L("tA")},IF({L("chosen_option")}="B",{L("tB")},{L("tC")})))))')
                elif k == "f_need":
                    c.value = f'=IF({L("in_stock")}="",{L("recommended_qty")},MAX({L("f_target")}-{L("in_stock")}-N({L("in_transit")}),0))'
                elif k == "f_diff":
                    c.value = f'=IF(OR({L("in_final")}="",{L("f_need")}=0),"",{L("in_final")}/{L("f_need")}-1)'
                elif k == "f_check":
                    c.value = (f'=IF({L("in_decision")}="","Chưa quyết định",'
                               f'IF(OR({L("in_reason")}="",{L("in_reviewer")}=""),"⚠ Thiếu lý do/người duyệt",'
                               f'IF(AND({L("in_final")}<>"",N({L("in_final")})+N({L("in_stock")})+N({L("in_transit")})>{L("tC")}*1.001,{L("in_stock")}<>""),"⚠ Vượt mục tiêu P90",'
                               f'IF(AND({L("f_diff")}<>"",ABS(N({L("f_diff")}))>0.3),"⚠ Lệch > 30% so với nhu cầu",'
                               f'IF({L("promo_extrapolation")}="Có","ℹ Kiểm tra ngoại suy KM","✓ Hợp lệ")))))')
            if fmt:
                c.number_format = fmt
            if k in ("reason_codes", "in_reason", "promotion_stance"):
                c.alignment = Alignment(wrap_text=True, vertical="top")
    # dropdown
    dv1 = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Giá trị không hợp lệ", error="Chọn: " + ", ".join(DECISIONS))
    dv2 = DataValidation(type="list", formula1='"' + ",".join(OPTIONS) + '"', allow_blank=True, showErrorMessage=True)
    dv3 = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, error="Nhập số ≥ 0")
    dv4 = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True, error="Nhập ngày hợp lệ")
    for dv in (dv1, dv2, dv3, dv4):
        ws.add_data_validation(dv)
    dv1.add(f"{col['in_decision']}{r1}:{col['in_decision']}{r2}"); dv2.add(f"{col['in_option']}{r1}:{col['in_option']}{r2}")
    for k in ("in_stock", "in_transit", "in_final"):
        dv3.add(f"{col[k]}{r1}:{col[k]}{r2}")
    dv4.add(f"{col['in_date']}{r1}:{col['in_date']}{r2}")
    # định dạng có điều kiện
    rng = f"{col['f_check']}{r1}:{col['f_check']}{r2}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col["f_check"]}{r1},1)="⚠"'], fill=PatternFill("solid", fgColor="F9DEDA"), font=Font(bold=True, color="9B2C22")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col["f_check"]}{r1},1)="✓"'], fill=PatternFill("solid", fgColor="DCEFE2"), font=Font(bold=True, color="276749")))
    for k, fc in (("HIGH", "F9DEDA"), ("MEDIUM", "FBEBD0"), ("LOW", "DCEFE2")):
        ws.conditional_formatting.add(f"{col['risk']}{r1}:{col['risk']}{r2}", FormulaRule(formula=[f'{col["risk"]}{r1}="{k}"'], fill=PatternFill("solid", fgColor=fc)))
    dec = f"{col['in_decision']}{r1}:{col['in_decision']}{r2}"
    for k, fc in (("Duyệt", "DCEFE2"), ("Điều chỉnh", "FBEBD0"), ("Từ chối", "F9DEDA"), ("Yêu cầu xem lại", "E2E8F0")):
        ws.conditional_formatting.add(dec, FormulaRule(formula=[f'{col["in_decision"]}{r1}="{k}"'], fill=PatternFill("solid", fgColor=fc)))
    ws.freeze_panes = ws.cell(r1, 5)
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(LAYOUT))}{r2}"
    ws.protection.sheet = True; ws.protection.autoFilter = False; ws.protection.sort = False
    ws.protection.formatColumns = False; ws.protection.formatRows = False
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{HR}:{HR}"

    # ------------------------------------------------ Tổng hợp
    wt = wb.create_sheet("Tổng hợp"); wt.sheet_view.showGridLines = False
    wt["A1"] = "Tổng hợp quyết định"; wt["A1"].font = Font(size=16, bold=True, color=INK)
    wt["A2"] = "Tự cập nhật khi nhập ở sheet 'Mẫu quyết định'."; wt["A2"].font = Font(color="5B6B70")
    S = "'Mẫu quyết định'!"
    dcol, fcol, ncol, kcol, rcol = (f"{S}${col[k]}${r1}:${col[k]}${r2}" for k in ("in_decision", "in_final", "f_need", "f_check", "risk"))
    hdr = ["Chỉ tiêu", "Giá trị"]
    for j, h in enumerate(hdr, 1):
        c = wt.cell(4, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR_LOCK
    items = [("Tổng số phiếu", f"=ROWS({dcol})", "0"), ("Đã quyết định", f"=COUNTA({dcol})", "0"), ("Chưa quyết định", "=B5-B6", "0"),
             ("Tỷ lệ hoàn thành", "=IF(B5=0,0,B6/B5)", "0.0%")] + [(f"Số phiếu: {d}", f'=COUNTIF({dcol},"{d}")', "0") for d in DECISIONS] + [
             ("Tổng SL cần theo tồn thực", f"=SUM({ncol})", "#,##0"), ("Tổng số lượng cuối đã duyệt", f'=SUMIFS({fcol},{dcol},"Duyệt")+SUMIFS({fcol},{dcol},"Điều chỉnh")', "#,##0"),
             ("Số dòng có cảnh báo ⚠", f'=COUNTIF({kcol},"⚠*")', "0"),
             ("Phiếu HIGH chưa quyết định", f'=COUNTIFS({rcol},"HIGH",{dcol},"")', "0")]
    for i, (k, f, fmt) in enumerate(items, 5):
        wt.cell(i, 1, k).border = THIN; c = wt.cell(i, 2, f); c.number_format = fmt; c.font = Font(bold=True, color=PETROL); c.border = THIN
    wt.column_dimensions["A"].width = 34; wt.column_dimensions["B"].width = 18
    wt["D4"] = "Ghi chú"; wt["D4"].font = Font(bold=True)
    wt["D5"] = f"Tạo lúc {datetime.now():%d/%m/%Y %H:%M} từ outputs/ của pipeline (Stage 05, 07, 08)."
    wt["D6"] = "Sau khi duyệt, có thể nhập lại quyết định vào dashboard (trang 07) để ghi nhật ký HITL."
    wb.active = 1
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()
